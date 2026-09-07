"""
Rebuild the shared workbook from the log, with the days in date order.

WHY THIS IS A SCRIPT AND NOT A BUTTON
The workbook is append-only on purpose: a new day is added at the bottom and no
column or row is ever inserted, so anything typed straight into Excel stays
where it was put. The price is that a day which was MISSED lands at the bottom
too. On 2026-09-03 twenty days had to be added at once (an interrupted first
run had left the sheet with three days out of twenty-three), and putting them
in date order means writing the whole sheet again. That is not something a
program should do to a file people edit by hand, so it is a deliberate,
run-by-hand job.

BEFORE RUNNING IT, prove nothing would be lost: every value in the sheet has to
be reproducible from spfe_log.csv. Run with --diff, read the list, and only
then --write. Cells that differ only because the field file now rounds them,
and cells the sheet never received, are fine; anything else is somebody's hand
typing and would be destroyed.

  python testing/rebuild_workbook.py --diff     what the log cannot reproduce
  python testing/rebuild_workbook.py            build it locally and check it
  python testing/rebuild_workbook.py --write    replace the shared workbook

The workbook is copied into %APPDATA%\\SPFE_Values first, and the replacement
is a move onto the target, so a failure leaves the old file untouched. Excel
must not have it open.
"""
import os
import shutil
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import spfe_store as store                                    # noqa: E402


def share_folder() -> Path:
    cfg = store.load_config()
    root, how = store.resolve_share(cfg)
    if not root:
        print("Neither share name answered.")
        raise SystemExit(1)
    folder = store.share_dir(root, cfg)
    print(f"share: {folder}  ({how})")
    return folder


def blocks(path: Path) -> "list[date]":
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    ws = wb[store.SHEET_VALUES] if store.SHEET_VALUES in wb.sheetnames \
        else wb.worksheets[0]
    out = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, datetime):
            v = v.date()
        if isinstance(v, date):
            out.append(v)
    return out


def records(folder: Path, fields):
    """Every moment the log holds, minus the deleted ones, in the order the
    columns of a block are meant to be in: Morning, At the end, then the extra
    moments by their time."""
    gone = set(store.read_json_map(folder / store.TOMBSTONES_NAME, "deleted"))
    out = []
    for row in store.read_csv(folder / store.CSV_NAME):
        if row.get("datetime", "") in gone:
            continue
        rec = store.record_from_csv(row, fields)
        if rec is not None:
            out.append(rec)
    order = {store.SLOT_MORNING: 0, store.SLOT_EVENING: 1}
    out.sort(key=lambda r: (r.day, order.get(r.slot, 2), r.when))
    return out


def show_diff(folder: Path, fields) -> None:
    """Every cell in the sheet that the log cannot reproduce.

    Reads the sheet in WHICHEVER layout it is in. Until 2026-09-03 a quantity
    had one cell per moment, holding "150 ; -250 ; 2500"; it has one cell per
    number now. A block is recognised by its second line: the new shape carries
    the names of the numbers there, the old one carries the first quantity.
    """
    from openpyxl import load_workbook

    by_day: dict = {}
    for rec in records(folder, fields):
        by_day.setdefault(rec.day, {})[rec.heading()] = rec

    wb = load_workbook(folder / store.WORKBOOK_NAME)
    ws = wb[store.SHEET_VALUES] if store.SHEET_VALUES in wb.sheetnames \
        else wb.worksheets[0]
    parts = store.parts_per_moment(fields)
    found = 0
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, datetime):
            v = v.date()
        if not isinstance(v, date):
            continue
        want = by_day.get(v, {})
        old = store._block_is_old_shape(ws, r, fields)
        step = 1 if old else parts
        first = r + (1 if old else store.BLOCK_HEAD_LINES)
        for m in range(2 + store.MAX_EXTRA_COLS):
            c = store.FIRST_VALUE_COL + m * step
            head = ws.cell(r, c).value
            if head in (None, ""):
                continue
            rec = want.get(str(head))
            for i, row_def in enumerate(fields.rows):
                if old or row_def.is_text or len(row_def.columns) == 1:
                    pairs = [(c, store.format_cell(row_def, rec.values)
                              if rec is not None else None)]
                else:
                    pairs = [
                        (c + k, store.format_value(row_def, column, rec.values)
                         if rec is not None else None)
                        for k, column in enumerate(row_def.columns)]
                for col, exp in pairs:
                    got = ws.cell(first + i, col).value
                    blank = (None, "", "-")
                    if got in blank and exp in blank:
                        continue
                    if str(got or "") == str(exp or ""):
                        continue
                    found += 1
                    print(f"  row {first + i:4d}  {head} / "
                          f"{row_def.label or '(blank)'}: "
                          f"sheet={got!r} log={exp!r}")
    print(f"{found} cell(s) the log cannot reproduce.")


def main() -> int:
    fields = store.load_fields()
    folder = share_folder()
    target = folder / store.WORKBOOK_NAME

    if "--diff" in sys.argv:
        show_diff(folder, fields)
        return 0

    recs = records(folder, fields)
    days = sorted({r.day for r in recs})
    if not recs:
        print("The log is empty - nothing to build.")
        return 1
    print(f"log: {len(recs)} moments over {len(days)} days "
          f"({days[0]} .. {days[-1]})")

    tmp = Path(tempfile.mkdtemp(prefix="spfe_rebuild_"))
    fresh = tmp / store.WORKBOOK_NAME
    entries = store.read_json_map(folder / store.CAMPAIGNS_NAME, "days")
    for rec in recs:
        store.write_workbook_record(
            fresh, fields, rec,
            campaign=store.campaign_for(entries, rec.day)[0])
    store.write_workbook_trends(fresh, fields, recs, log_fn=print)
    print(f"built {fresh}  {fresh.stat().st_size} B")

    seen = blocks(fresh)
    print(f"blocks: {len(seen)}, in date order: {seen == sorted(seen)}")
    if seen != sorted(seen) or len(seen) != len(days):
        print("REFUSING: the rebuilt sheet is not what it should be.")
        return 1

    if "--write" not in sys.argv:
        print("\nDry run. Add --write to replace the shared workbook.")
        return 0

    backup = (store.local_dir()
              / f"SPFE values.before-rebuild-{datetime.now():%Y%m%d-%H%M%S}.xlsx")
    if target.exists():
        shutil.copy2(target, backup)
        print(f"backed up the shared workbook to\n  {backup}")

    staged = target.with_name(target.stem + ".rebuilding.xlsx")
    shutil.copy2(fresh, staged)
    try:
        os.replace(staged, target)
    except PermissionError as exc:
        os.remove(staged)
        print(f"REFUSED: {target.name} is open in Excel ({exc}). "
              "Nothing was changed.")
        return 1
    print(f"replaced {target}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
